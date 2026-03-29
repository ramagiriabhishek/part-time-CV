#!/usr/bin/env python3
"""Convert tailored CV markdown files to simple UK-style Word documents."""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.shared import Pt


def add_runs_with_bold(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part:
            paragraph.add_run(part)


def md_to_docx(md_path: Path, out_path: Path) -> None:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    style.paragraph_format.space_after = Pt(4)

    lines = md_path.read_text(encoding="utf-8").splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip():
            i += 1
            continue
        if line.strip() == "---":
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line[2:].strip())
            run.bold = True
            run.font.size = Pt(18)
            run.font.name = "Calibri"
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs_with_bold(p, line[2:].strip())
            i += 1
            continue

        p = doc.add_paragraph()
        add_runs_with_bold(p, line.strip())
        i += 1

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def main() -> None:
    base = Path(__file__).resolve().parent
    files = [
        "CV-01-Crew-Chef-Team-Member.md",
        "CV-02-Warehouse-Operative-Amazon-FC.md",
        "CV-03-Customer-Service-Assistant.md",
        "CV-04-Hospitality-Assistant.md",
    ]
    out_dir = base / "word"
    for name in files:
        src = base / name
        if not src.exists():
            print(f"Missing {src}", file=sys.stderr)
            sys.exit(1)
        stem = src.stem
        md_to_docx(src, out_dir / f"{stem}.docx")
        print(f"Wrote {out_dir / f'{stem}.docx'}")


if __name__ == "__main__":
    main()
